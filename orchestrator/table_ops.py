"""Merge-cell-safe helpers for Word (.docx) and Excel (.xlsx) templates.

Key correctness properties this module upholds
---------------------------------------------
1.  Writing a Word cell uses ``run.text`` mutation instead of ``cell.text = ...``
    so paragraph / run formatting (font, size, alignment) is preserved.
2.  Merged cells never get written twice: for each logical coordinate we map to
    the underlying anchor ``<w:tc>`` (python-docx returns the anchor for any
    slave coordinate). A set of anchor ids deduplicates writes.
3.  Excel merged slave cells are detected via ``openpyxl.cell.cell.MergedCell``
    and skipped silently (writing to them would raise).
4.  For Word templates the paragraph preceding the first table is used as
    ``context`` so the LLM can learn "this table is about 德州市 / 2024 Q1".
    We walk ``doc.element.body`` in XML order so the context is tied to the
    nearest structural neighbour of the table, not to the whole document.

The coordinate convention inside the orchestrator is **0-based (r, c)**
encoded as ``"r_c"``. For Excel that translates to 1-based row/col at the
openpyxl boundary only.
"""
from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Public types
# ---------------------------------------------------------------------------

@dataclass
class TemplateInfo:
    kind: str              # "word" | "excel"
    doc_obj: Any           # docx.Document or openpyxl.Workbook
    headers: List[str]     # flattened first-row headers (merged continuations -> "")
    context: str           # nearest paragraph(s) preceding the first table (word only)
    table_index: int = 0   # which table inside the doc (Word multi-table; Excel always 0)


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def _docx_paragraphs_before_first_table(doc) -> List[str]:
    """Collect non-empty paragraphs that appear before the first <w:tbl>."""
    return _docx_context_for_table(doc, 0)


def _docx_context_for_table(doc, table_index: int) -> List[str]:
    """Collect non-empty paragraphs immediately preceding the *table_index*-th table.

    For table 0 this is every paragraph before the first ``<w:tbl>``.  For
    table N (N>0) we collect paragraphs between table N-1 and table N.
    """
    from docx.oxml.ns import qn

    w_t = qn("w:t")
    tables_seen = 0
    paragraphs: List[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            if tables_seen == table_index:
                break  # stop: we've reached the target table
            tables_seen += 1
            paragraphs.clear()  # reset: previous paragraphs belong to an earlier table
        elif tag == "p":
            texts = child.findall(".//" + w_t)
            text = "".join(t.text or "" for t in texts).strip()
            if text:
                paragraphs.append(text)
    return paragraphs


def _docx_header_row(table) -> List[str]:
    """Read the first row, marking merged continuations as empty string.

    Dedup key is the tc element itself (not ``id(tc)``): Python can recycle an
    integer id once the underlying proxy is GC'd mid-loop, which on a longer
    table would silently collapse unrelated cells. The anchor set here holds
    the element directly so lxml's proxy cache stays warm.
    """
    header_cells = table.rows[0].cells
    seen_tcs: set = set()
    headers: List[str] = []
    for cell in header_cells:
        tc = cell._tc
        if tc in seen_tcs:
            headers.append("")
        else:
            seen_tcs.add(tc)
            headers.append(cell.text.strip())
    return headers


def load_word_template(path: str) -> List[TemplateInfo]:
    """Return one :class:`TemplateInfo` **per table** found in the Word doc.

    Each entry carries its own ``headers`` (from the table's first row) and
    ``context`` (paragraphs between the previous table and this one), so the
    caller can process multi-table templates where every table has a different
    scope (e.g. different city or time-period).
    """
    from docx import Document

    doc = Document(path)
    if not doc.tables:
        raise ValueError("Word template contains no table")
    infos: List[TemplateInfo] = []
    for idx, tbl in enumerate(doc.tables):
        ctx_lines = _docx_context_for_table(doc, idx)
        infos.append(TemplateInfo(
            kind="word",
            doc_obj=doc,
            headers=_docx_header_row(tbl),
            context="\n".join(ctx_lines),
            table_index=idx,
        ))
    logger.info("load_word_template: %d table(s) in %s", len(infos), path)
    return infos


def _set_cell_text_preserve_style(cell, text: str) -> None:
    """Write ``text`` into a Word cell without blowing away formatting.

    ``cell.text = "x"`` in python-docx wipes all paragraphs and runs, losing
    font, bold, colour, etc. We instead mutate the first run's text and clear
    any following runs / paragraphs.
    """
    value = "" if text is None else str(text)
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(value)
        return
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(value)
    # remove trailing paragraphs inside the cell so we don't accumulate blanks
    for extra in paragraphs[1:]:
        el = extra._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _ensure_word_rows(table, required_rows: int) -> None:
    """Append blank rows until the table has at least ``required_rows``."""
    while len(table.rows) < required_rows:
        table.add_row()


def get_word_grid(doc, table_index: int = 0) -> List[List[str]]:
    """Return the *table_index*-th table as a 2-D list of strings.

    Merged continuation cells render as ``""`` so that ``grid[r][c]`` stays
    column-aligned with ``headers``.

    See :func:`_docx_header_row` for why dedup uses tc identity rather than
    ``id(tc)`` (id recycling after GC).
    """
    table = doc.tables[table_index]
    grid: List[List[str]] = []
    for row in table.rows:
        seen_tcs: set = set()
        row_vals: List[str] = []
        for cell in row.cells:
            tc = cell._tc
            if tc in seen_tcs:
                row_vals.append("")
            else:
                seen_tcs.add(tc)
                row_vals.append(cell.text.strip())
        grid.append(row_vals)
    return grid


def write_word_cells(doc, data: Dict[str, str], table_index: int = 0) -> None:
    """Write ``{"r_c": value}`` pairs into the *table_index*-th Word table.

    - Rows are appended as needed so callers may address coordinates beyond the
      template's starting row count.
    - Each anchor cell is written at most once even if several coordinates
      resolve to it (merged regions).

    IMPORTANT: dedup key is the tc element itself, not ``id(tc)``. Using the
    integer id leaves the set with only weak information about the proxy;
    once Python GCs the previous iteration's cell wrapper the lxml proxy
    cache can drop the tc, its Python id gets reused for a later unrelated
    cell, and we end up silently skipping a legitimate write. S3 integration
    testing saw random cells come out blank because of exactly this.
    """
    if not data:
        return
    table = doc.tables[table_index]
    max_row = max(int(k.split("_")[0]) for k in data.keys())
    _ensure_word_rows(table, max_row + 1)

    written_tcs: set = set()
    for coord, value in data.items():
        try:
            r, c = (int(x) for x in coord.split("_"))
        except ValueError:
            logger.warning("invalid coord %r, skipped", coord)
            continue
        try:
            cell = table.cell(r, c)
        except IndexError:
            logger.warning("cell %s out of bounds in Word table", coord)
            continue
        tc = cell._tc
        if tc in written_tcs:
            # another coordinate already wrote to the merged anchor
            continue
        written_tcs.add(tc)  # set holds a strong ref so the proxy sticks
        _set_cell_text_preserve_style(cell, value)


# ---------------------------------------------------------------------------
# Excel (.xlsx)
# ---------------------------------------------------------------------------

def load_excel_template(path: str) -> TemplateInfo:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers: List[str] = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = ["" if v is None else str(v).strip() for v in row]
        break
    return TemplateInfo(kind="excel", doc_obj=wb, headers=headers, context="")


def get_excel_grid(wb) -> List[List[str]]:
    ws = wb.active
    grid: List[List[str]] = []
    for row in ws.iter_rows(values_only=True):
        grid.append(["" if v is None else str(v).strip() for v in row])
    return grid


def _is_excel_merged_slave(ws, row_1based: int, col_1based: int) -> bool:
    import openpyxl

    cell = ws.cell(row=row_1based, column=col_1based)
    return isinstance(cell, openpyxl.cell.cell.MergedCell)


def write_excel_cells(wb, data: Dict[str, str]) -> None:
    if not data:
        return
    ws = wb.active
    for coord, value in data.items():
        try:
            r, c = (int(x) for x in coord.split("_"))
        except ValueError:
            logger.warning("invalid coord %r, skipped", coord)
            continue
        xr, xc = r + 1, c + 1  # openpyxl uses 1-based indices
        if _is_excel_merged_slave(ws, xr, xc):
            logger.debug("skip merged-slave cell %s", coord)
            continue
        ws.cell(row=xr, column=xc, value="" if value is None else str(value))


# ---------------------------------------------------------------------------
# Common
# ---------------------------------------------------------------------------

def classify_template(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".docx", ".doc", ".docm"):
        return "word"
    if ext in (".xlsx", ".xls", ".xlsm"):
        return "excel"
    raise ValueError(f"unsupported template extension: {ext}")


def save_template(info: TemplateInfo, out_path: str) -> None:
    info.doc_obj.save(out_path)
